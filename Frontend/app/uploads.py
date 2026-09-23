"""Bounded extraction; documents are never saved or executed."""
import base64
import io
import zipfile
from pathlib import Path

MAX_BYTES = 5 * 1024 * 1024


def extract(name, data):
    if not data or len(data)>MAX_BYTES:
        raise ValueError('Файл бос немесе көлемі 5 MB-тан асады.')
    ext = Path(name).suffix.lower()
    if ext in ['.docx','.xlsx']:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > 20*1024*1024 or len(z.infolist()) > 2000:
                raise ValueError('Құжаттың ашылған көлемі тым үлкен.')
    if ext == '.docx':
        from docx import Document
        doc = Document(io.BytesIO(data))
        text = '\n'.join(p.text for p in doc.paragraphs)
        text += '\n'+'\n'.join(' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
    elif ext == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data),read_only=True,data_only=True)
        try:
            lines=[]
            for sheet in wb.worksheets[:5]:
                for row in sheet.iter_rows(max_row=300,max_col=20,values_only=True):
                    lines.append(' | '.join(str(x) for x in row if x is not None))
            text='\n'.join(lines)
        finally:
            wb.close()
    elif ext == '.pdf':
        from pypdf import PdfReader
        pdf = PdfReader(io.BytesIO(data))
        if pdf.is_encrypted or len(pdf.pages)>30:
            raise ValueError('Құпиясөзсіз, 30 бетке дейінгі PDF жіберіңіз.')
        text='\n'.join(page.extract_text() or '' for page in pdf.pages)
        if not text.strip():
            raise ValueError('PDF ішінде мәтін табылмады. Сканды JPEG ретінде жіберіңіз немесе артикулды жазыңыз.')
    elif ext in ['.jpg','.jpeg','.png']:
        from PIL import Image, ImageOps
        Image.MAX_IMAGE_PIXELS = 16000000
        with Image.open(io.BytesIO(data)) as im:
            im = ImageOps.exif_transpose(im).convert('RGB')
            im.thumbnail((1600,1600))
            out=io.BytesIO(); im.save(out,format='JPEG',quality=85)
        return dict(name=name, image='data:image/jpeg;base64,'+base64.b64encode(out.getvalue()).decode())
    else:
        raise ValueError('Қолдау: .xlsx, .docx, .pdf, .jpeg, .jpg, .png. Ескі .xls/.doc файлдарын жаңа форматқа сақтаңыз.')
    if not text.strip():
        raise ValueError('Құжаттан мәтін табылмады.')
    return dict(name=name,text=text[:20000],truncated=len(text)>20000)
