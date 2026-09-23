import io
import json
import time
import pytest
from fastapi.testclient import TestClient
from app import store, agent, uploads
from app.main import app


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(store,'DB_PATH',tmp_path/'test.sqlite3')
    monkeypatch.delenv('OPENAI_API_KEY',raising=False)
    with TestClient(app,base_url='http://localhost') as c:
        session=c.get('/api/session').json()
        c.headers['X-CSRF-Token']=session['csrf']
        yield c


def post_chat(c,message,**extra):
    return c.post('/api/chat',data=dict(message=message,mode='demo',**extra))


def proposal(c,quantity=2,pid='p002'):
    r=c.post('/api/cart/prepare',json=dict(product_id=pid,quantity=quantity))
    assert r.status_code==200,r.text
    return r.json()


def accept(c,p):
    return c.post('/api/cart/confirm',json=dict(token=p['token'],accepted=True))


def test_catalog_facts_certificate(client):
    result=post_chat(client,'DEMO-002 бар ма?').json()
    p=result['products'][0]
    assert p['article']=='DEMO-002' and p['price']==2660
    assert store.stock(p)==14
    assert client.get(p['certificate_url']).status_code==200
    assert 'ТЕСТІЛІК' in client.get(p['certificate_url']).text


def test_zero_stock_analogs(client):
    result=post_chat(client,'DEMO-001 бар ма?').json()
    assert len(result['products'])>=2
    assert 'аналог' in result['answer']
    assert any(p.get('analog_reason') for p in result['products'])
    assert all(p['specifications']==result['products'][0]['specifications'] for p in result['products'])


def test_terms(client):
    answer=post_chat(client,'Төлем және жеткізу шарттары').json()['answer']
    assert all(word in answer for word in ['төлем','жеткізу','ең аз партия','Тест'])


def test_explicit_confirmation_and_replay(client):
    p=proposal(client)
    assert client.get('/api/cart').json()['count']==0
    assert client.post('/api/cart/confirm',json=dict(token=p['token'],accepted=False)).status_code==400
    assert accept(client,p).json()['count']==2
    assert accept(client,p).status_code==400
    cart=client.get('/api/cart').json()
    assert cart['count']==2 and client.get(cart['url']).status_code==200


def test_plain_yes_cannot_add(client):
    proposal(client)
    post_chat(client,'иә')
    assert client.get('/api/cart').json()['count']==0
    assert post_chat(client,'иә, қос').status_code==400


def test_text_confirmation(client):
    r=post_chat(client,'DEMO-002 тауарынан 2 дана қос')
    assert r.status_code==200,r.text
    p=r.json()['pending']
    assert client.get('/api/cart').json()['count']==0
    result=post_chat(client,'иә, қос',confirmation_token=p['token'])
    assert result.json()['cart']['count']==2


@pytest.mark.parametrize('quantity',[0,-1,1.5,True,'2',10001])
def test_bad_quantities(client,quantity):
    assert client.post('/api/cart/prepare',json=dict(product_id='p002',quantity=quantity)).status_code==422


def test_stock_limit_including_existing_cart(client):
    assert client.post('/api/cart/prepare',json=dict(product_id='p002',quantity=15)).status_code==400
    p=proposal(client,12);assert accept(client,p).status_code==200
    assert client.post('/api/cart/prepare',json=dict(product_id='p002',quantity=3)).status_code==400


@pytest.mark.parametrize('change',['stock','price','expiry'])
def test_revalidation_at_confirmation(client,change):
    p=proposal(client)
    with store.db() as c:
        if change=='expiry':
            c.execute('UPDATE pending SET created=?',(time.time()-601,))
        else:
            product=store.product('p002')
            if change=='stock':
                for warehouse in product['stock']:warehouse['quantity']=0
            else:product['price']+=10
            c.execute('UPDATE products SET payload=? WHERE id=?',(json.dumps(product),'p002'))
    assert accept(client,p).status_code==400
    assert client.get('/api/cart').json()['count']==0


def test_session_isolation_and_csrf(client):
    p=proposal(client)
    with TestClient(app,base_url='http://localhost') as other:
        token=other.get('/api/session').json()['csrf']
        assert other.post('/api/cart/confirm',json=dict(token=p['token'],accepted=True)).status_code==403
        other.headers['X-CSRF-Token']=token
        assert accept(other,p).status_code==400
        assert other.get('/api/cart').json()['count']==0
    assert client.post('/api/cart/cancel',json={},headers={'Origin':'https://evil.example'}).status_code==403


def test_cancel_and_stale_proposal(client):
    p=proposal(client)
    newer=proposal(client,3)
    assert accept(client,p).status_code==400
    client.post('/api/cart/cancel',json={})
    assert accept(client,newer).status_code==400


def test_docx_upload(client):
    from docx import Document
    doc=Document();doc.add_paragraph('DEMO-002');buffer=io.BytesIO();doc.save(buffer)
    r=client.post('/api/chat',data={'message':'Файлдағы тауарды тап','mode':'demo'},files={'file':('spec.docx',buffer.getvalue())})
    assert r.status_code==200,r.text
    assert r.json()['products'][0]['article']=='DEMO-002'


def test_xlsx_extraction(client):
    from openpyxl import Workbook
    wb=Workbook();wb.active.append(['DEMO-002',2]);buffer=io.BytesIO();wb.save(buffer)
    assert 'DEMO-002' in uploads.extract('spec.xlsx',buffer.getvalue())['text']


def test_pdf_and_image_extraction(client):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject,NameObject,DictionaryObject
    writer=PdfWriter();page=writer.add_blank_page(width=200,height=200)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
    stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 10 100 Td (DEMO-002) Tj ET')
    page[NameObject('/Contents')]=writer._add_object(stream)
    buffer=io.BytesIO();writer.write(buffer)
    assert 'DEMO-002' in uploads.extract('spec.pdf',buffer.getvalue())['text']
    from PIL import Image
    buffer=io.BytesIO();Image.new('RGB',(100,100)).save(buffer,format='JPEG')
    assert uploads.extract('photo.jpg',buffer.getvalue())['image'].startswith('data:image/jpeg;base64,')


def test_bad_upload(client):
    r=client.post('/api/chat',data={'message':'Тауарды тап'},files={'file':('spec.exe',b'bad')})
    assert r.status_code==400
    with pytest.raises(ValueError):uploads.extract('huge.pdf',b'a'*(uploads.MAX_BYTES+1))


@pytest.mark.parametrize('quantity',['-2','1.5','1,5'])
def test_invalid_quantity_in_chat(client,quantity):
    assert post_chat(client,f'DEMO-002 {quantity} дана қос').status_code==400
    assert client.get('/api/cart').json()['count']==0


def test_payment_data_not_stored(client):
    assert post_chat(client,'4111 1111 1111 1111').status_code==400
    assert not store.history(client.cookies['ekt_session'])


def test_technical_filters(client):
    assert client.get('/api/products?q=B16').json()==[]
    assert client.get('/api/products?q=C16+2P').json()==[]
    products=client.get('/api/products?q=кабель+3x2.5').json()
    assert products and all(p['specifications']['cross_section_mm2']==2.5 for p in products)


def test_cloud_requires_key_and_consent(client,monkeypatch):
    assert client.post('/api/chat',data={'message':'Hi','mode':'ai','cloud_consent':'true'}).status_code==400
    monkeypatch.setenv('OPENAI_API_KEY','test-key')
    assert client.post('/api/chat',data={'message':'Hi','mode':'ai'}).status_code==400


def test_ai_tool_loop_with_mocked_transport(client,monkeypatch):
    # No paid API call: verify Responses continuation and tool output contracts.
    captured=[]
    class Response:
        def __init__(self,data):self.data=data
        def raise_for_status(self):pass
        def json(self):return self.data
    class FakeClient:
        def __init__(self,**kwargs):pass
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def post(self,url,**kwargs):
            captured.append(json.loads(json.dumps(kwargs['json'])))
            if len(captured)==1:
                return Response({'status':'completed','output':[{'type':'function_call','name':'get_product','arguments':'{"product_id":"p002"}','call_id':'call-1'}]})
            return Response({'status':'completed','output':[{'type':'message','content':[{'type':'output_text','text':'DEMO-002: 14 дана, 2660 ₸.'}]}]})
    monkeypatch.setenv('OPENAI_API_KEY','test-key')
    sid=client.cookies['ekt_session']
    monkeypatch.setattr(agent.httpx,'Client',FakeClient)
    answer,products,mode=agent.cloud(sid,'DEMO-002 бар ма?')
    assert mode=='ai' and products[0]['id']=='p002'
    assert captured[1]['input'][-1]['type']=='function_call_output'
    assert not any(t['name']=='confirm_cart_add' for t in captured[0]['tools'])
